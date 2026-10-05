"""excel_tracker.py - quản lý và cập nhật results.xlsx đúng chuẩn GUIDE.md mục 6.1.
Tạo sẵn 7 sheet với format đẹp, tự động ghi nhận số liệu từ các thí nghiệm.
"""
from __future__ import annotations

import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from pathlib import Path
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_PATH = Path("submissions/2A202602368_duong_van_thanh/results.xlsx")


def get_header_style():
    fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return fill, font, align


def init_excel(path: Path = EXCEL_PATH) -> None:
    """Khởi tạo file results.xlsx với 7 sheet chuẩn."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    # Xoá sheet mặc định
    wb.remove(wb.active)

    fill, font, align = get_header_style()

    sheets_headers = {
        "Backbones": [
            "exp_id", "backbone", "tag", "params_m", "gmacs", "img_size",
            "epochs", "seed", "val_macro_f1", "val_top1", "train_sec_per_epoch",
            "latency_p50_ms", "latency_p95_ms", "notes"
        ],
        "Training": [
            "exp_id", "backbone", "axis", "diff_from_t00", "seed",
            "val_macro_f1", "val_top1", "delta_vs_t00", "f1_chinee_apple", "f1_snake_weed", "notes"
        ],
        "Inference": [
            "exp_id", "method", "model_ckpt", "k_views", "val_macro_f1", "val_top1",
            "val_ece", "latency_p50_ms", "latency_p95_ms", "latency_p99_ms", "throughput_img_s", "rel_cost"
        ],
        "Final": [
            "exp_id", "config", "seed", "val_macro_f1", "test_macro_f1", "test_top1", "test_ece", "notes"
        ],
        "PerClass": [
            "class_id", "class_name", "test_samples", "precision", "recall", "f1_score", "model_tag"
        ],
        "Latency": [
            "config", "gpu", "dtype", "batch", "fuse_bn", "p50_ms", "p95_ms", "p99_ms", "img_per_s"
        ],
        "Summary": [
            "rank", "exp_id", "type", "description", "val_macro_f1", "val_top1", "latency_p95_ms", "notes"
        ],
    }

    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for sheet_name, headers in sheets_headers.items():
        ws = wb.create_sheet(title=sheet_name)
        ws.row_dimensions[1].height = 26
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = fill
            cell.font = font
            cell.alignment = align
            ws.column_dimensions[get_column_letter(col_idx)].width = max(len(h) + 4, 14)
        ws.freeze_panes = "A2"

    wb.save(path)
    print(f"Khởi tạo thành công {path} với {len(sheets_headers)} sheets!")


def add_row(sheet_name: str, row_data: List[Any], path: Path = EXCEL_PATH) -> None:
    """Thêm một dòng dữ liệu vào sheet tương ứng (hoặc cập nhật nếu ID đã tồn tại)."""
    if not path.exists():
        init_excel(path)
    wb = openpyxl.load_workbook(path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet {sheet_name} không tồn tại trong {path}")
    ws = wb[sheet_name]

    # Kiểm tra xem dòng với ID này đã tồn tại chưa (ở cột A hoặc B)
    # Với Summary: exp_id ở cột B (col 2). Với các sheet khác: exp_id ở cột A (col 1).
    id_col = 2 if sheet_name == "Summary" else 1
    target_id = str(row_data[id_col - 1]) if len(row_data) >= id_col else None

    existing_row = None
    if target_id is not None:
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=id_col).value) == target_id:
                existing_row = r
                break

    row_idx = existing_row if existing_row is not None else ws.max_row + 1
    font_data = Font(name="Segoe UI", size=10)

    for col_idx, val in enumerate(row_data, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = font_data
        if isinstance(val, float):
            cell.number_format = "0.0000" if val < 10 else "0.00"

    wb.save(path)


if __name__ == "__main__":
    init_excel()
